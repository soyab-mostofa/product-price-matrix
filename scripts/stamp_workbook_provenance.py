"""Stamp each local SKU with the workbook sheet and Excel row it came from.

Writes ``source_sheet`` and ``source_row`` into every product in
``verified_marketplace_research.json`` so a figure on the dashboard can be
walked back to a cell. ``source_row`` is the row as Excel numbers it (headers
are row 1, data starts at row 2), not an array index, so it can be typed
straight into the Name Box.

Keyed on (name, size) because the catalog holds same-name SKUs that differ
only by pack size — keying on name alone points two SKUs at one row.

Idempotent: re-running produces the same file.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Roopelle.com Final Excel Sheet.xlsx"
RESEARCH = ROOT / "verified_marketplace_research.json"

LOCAL_SHEET = "Local product "
ORGAGENIC_SHEET = "local product Orgagenic"


def name_key(value: object) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value or "").lower())).strip()


def size_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def workbook_index() -> tuple[dict[tuple[str, str], tuple[str, int]], dict[str, list[tuple[str, int]]]]:
    """(name, size) -> (sheet, excel_row), plus a name-only fallback index."""
    book = openpyxl.load_workbook(WORKBOOK, data_only=True)
    exact: dict[tuple[str, str], tuple[str, int]] = {}
    by_name: dict[str, list[tuple[str, int]]] = {}

    # (sheet, name column, size column)
    for sheet_name, name_col, size_col in (
        (LOCAL_SHEET, 1, 3),
        (ORGAGENIC_SHEET, 1, 2),
    ):
        sheet = book[sheet_name]
        for excel_row in range(2, sheet.max_row + 1):
            raw_name = sheet.cell(excel_row, name_col).value
            if raw_name is None or not str(raw_name).strip():
                continue
            key_name = name_key(raw_name)
            location = (sheet_name, excel_row)
            exact.setdefault((key_name, size_key(sheet.cell(excel_row, size_col).value)), location)
            by_name.setdefault(key_name, []).append(location)

    return exact, by_name


def main() -> int:
    exact, by_name = workbook_index()
    research = json.loads(RESEARCH.read_text(encoding="utf-8"))
    products = research.get("products", [])

    stamped = 0
    unresolved: list[str] = []

    for product in products:
        key_name = name_key(product.get("product_name"))
        location = exact.get((key_name, size_key(product.get("size"))))
        if location is None:
            # A workbook row can leave the size cell blank while the catalog
            # carries one; accept a name-only hit when it is unambiguous.
            candidates = by_name.get(key_name, [])
            location = candidates[0] if len(candidates) == 1 else None
        if location is None:
            unresolved.append(str(product.get("product_name")))
            continue
        sheet_name, excel_row = location
        product["source_sheet"] = sheet_name
        product["source_row"] = excel_row
        stamped += 1

    RESEARCH.write_text(
        json.dumps(research, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(f"products scanned  : {len(products)}")
    print(f"provenance stamped: {stamped}")
    print(f"unresolved        : {len(unresolved)}")
    for name in unresolved:
        print(f"  unresolved: {name}")
    return 1 if unresolved else 0


if __name__ == "__main__":
    sys.exit(main())
